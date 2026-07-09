# -*- coding: utf-8 -*-
"""Excel Ruidai の「完成 Project*_問題解答.xlsx」から類題用 JSON / config を生成する。"""

from __future__ import annotations

import json
import re
import shutil
import sys
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("openpyxl が必要です: pip install openpyxl", file=sys.stderr)
    sys.exit(1)

REPO_ROOT = Path(__file__).resolve().parents[2]
RUIDAI_DIR = REPO_ROOT / "MOSapp" / "Excel Ruidai"
JSON_OUT_DIR = REPO_ROOT / "MOSapp" / "mos_xaml_app" / "References" / "JSON"
CONFIG_PATH = REPO_ROOT / "MOSapp" / "mos_xaml_app" / "Assets" / "config.json"
CSV_OUT_DIR = REPO_ROOT / "MOSapp" / "mos_xaml_app" / "References" / "CSV"

# 演習 Tab1: 画面上の projectId → 教材 Checker 名（config.json と同じ）
TEXTBOOK_LIBRARY = {
    1: "ExcelChecker1_2",
    2: "ExcelChecker1_3",
    3: "ExcelChecker1_1",
    4: "ExcelChecker1_4",
    5: "ExcelChecker1_5",
    6: "ExcelChecker1_6",
    7: "ExcelChecker1_9",
    8: "ExcelChecker1_7",
    9: "ExcelChecker1_8",
    10: "ExcelChecker1_10",
}

HEADER_MARKERS = {"科目", "プロジェクトID", "タスクID", "問題文"}


def find_excel_for_project_set(project_id: int, set_no: int) -> Path | None:
    """Project フォルダ内の類題 xlsx をセット番号で検索。"""
    project_dir = RUIDAI_DIR / f"Project{project_id}"
    if not project_dir.is_dir():
        return None

    pattern = re.compile(rf"セット{set_no}[_\s]")
    candidates = []
    for p in project_dir.glob("*.xlsx"):
        if pattern.search(p.name):
            candidates.append(p)

    if not candidates:
        return None
    # 同名が複数あればファイル名でソートして先頭
    return sorted(candidates, key=lambda x: x.name)[0]


def parse_task_id(raw, fallback: int) -> int:
    if raw is None:
        return fallback
    text = str(raw).strip()
    if not text:
        return fallback
    if "-" in text:
        tail = text.split("-")[-1]
        if tail.isdigit():
            return int(tail)
    if text.isdigit():
        return int(text)
    return fallback


def parse_sheet_name(name: str) -> tuple[int, int] | None:
    m = re.match(r"^(\d+)-(\d+)$", str(name).strip())
    if not m:
        return None
    return int(m.group(1)), int(m.group(2))


def read_qa_workbook(path: Path) -> dict[int, dict[int, list[dict]]]:
    """{variant_set: {project_id: [tasks]}} を返す。"""
    wb = openpyxl.load_workbook(str(path), data_only=True)
    result: dict[int, dict[int, list[dict]]] = {}

    for sheet_name in wb.sheetnames:
        parsed = parse_sheet_name(sheet_name)
        if not parsed:
            print(f"  skip sheet (name): {path.name} / {sheet_name}")
            continue
        sheet_project, variant_set = parsed
        ws = wb[sheet_name]
        tasks: list[dict] = []
        task_seq = 0

        for row in range(1, ws.max_row + 1):
            c1 = ws.cell(row, 1).value
            c2 = ws.cell(row, 2).value
            c3 = ws.cell(row, 3).value
            c4 = ws.cell(row, 4).value
            c5 = ws.cell(row, 5).value

            if c1 is not None and str(c1).strip() in HEADER_MARKERS:
                continue
            if c4 is None or not str(c4).strip():
                continue

            project_id = int(c2) if c2 is not None and str(c2).strip().isdigit() else sheet_project
            task_seq += 1
            task_id = parse_task_id(c3, task_seq)
            description = str(c4).strip()
            answer = str(c5).strip() if c5 is not None else ""

            tasks.append(
                {
                    "taskId": task_id,
                    "description": description,
                    "_answer": answer,
                    "_projectId": project_id,
                }
            )

        if not tasks:
            print(f"  warn: no tasks in {path.name} / {sheet_name}")
            continue

        result.setdefault(variant_set, {}).setdefault(sheet_project, [])
        # 同一シート内で projectId が混在する場合は sheet の project を優先
        if sheet_project in result[variant_set] and result[variant_set][sheet_project]:
            print(f"  warn: duplicate sheet project {sheet_project} set {variant_set} in {path.name}")
        result[variant_set][sheet_project] = [
            {"taskId": t["taskId"], "description": t["description"]} for t in tasks
        ]

    return result


def merge_variant_data(all_data: list[dict[int, dict[int, list[dict]]]]) -> dict[int, dict[int, list[dict]]]:
    merged: dict[int, dict[int, list[dict]]] = {}
    for data in all_data:
        for set_no, projects in data.items():
            merged.setdefault(set_no, {})
            for project_id, tasks in projects.items():
                merged[set_no][project_id] = tasks
    return merged


def write_variant_json_files(merged: dict[int, dict[int, list[dict]]]) -> None:
    JSON_OUT_DIR.mkdir(parents=True, exist_ok=True)
    for set_no in range(1, 6):
        projects = merged.get(set_no, {})
        payload = {
            "projects": [
                {
                    "projectId": pid,
                    "tasks": sorted(tasks, key=lambda t: t["taskId"]),
                }
                for pid, tasks in sorted(projects.items())
            ]
        }
        out_path = JSON_OUT_DIR / f"MOS演習問題文一覧_PracticeVariant{set_no}.json"
        with out_path.open("w", encoding="utf-8") as f:
            json.dump(payload, f, ensure_ascii=False, indent=2)
        print(f"wrote {out_path.name} ({len(projects)} projects)")


def write_answer_csv(merged: dict[int, dict[int, list[dict]]], qa_paths: list[Path]) -> None:
    """解答手順付き CSV（社内原稿用。アプリは JSON のみ表示）。"""
    CSV_OUT_DIR.mkdir(parents=True, exist_ok=True)
    lines = ["グループ,プロジェクト,類題セット,タスク番号,問題文,解答操作"]

    # 再読込して解答列も含める
    all_rows = []
    for path in qa_paths:
        wb = openpyxl.load_workbook(str(path), data_only=True)
        for sheet_name in wb.sheetnames:
            parsed = parse_sheet_name(sheet_name)
            if not parsed:
                continue
            sheet_project, variant_set = parsed
            ws = wb[sheet_name]
            task_seq = 0
            for row in range(1, ws.max_row + 1):
                c1 = ws.cell(row, 1).value
                c2 = ws.cell(row, 2).value
                c3 = ws.cell(row, 3).value
                c4 = ws.cell(row, 4).value
                c5 = ws.cell(row, 5).value
                if c1 is not None and str(c1).strip() in HEADER_MARKERS:
                    continue
                if c4 is None or not str(c4).strip():
                    continue
                project_id = int(c2) if c2 is not None and str(c2).strip().isdigit() else sheet_project
                task_seq += 1
                task_id = parse_task_id(c3, task_seq)
                desc = str(c4).strip().replace('"', '""')
                ans = (str(c5).strip() if c5 is not None else "").replace('"', '""')
                all_rows.append(
                    f'1,{project_id},{variant_set},{task_id},"{desc}","{ans}"'
                )

    out = CSV_OUT_DIR / "MOS演習類題_問題解答一覧.csv"
    with out.open("w", encoding="utf-8-sig", newline="") as f:
        f.write("\n".join([lines[0]] + all_rows))
    print(f"wrote {out.name} ({len(all_rows)} rows)")


def sync_excel_files_to_mostest(merged: dict[int, dict[int, list[dict]]]) -> int:
    """Excel Ruidai → C:\\MOSTest\\Excel365\\Tab1\\PracticeVariant{n}\\project{m}.xlsx"""
    group = "1"
    copied = 0
    for set_no, projects in merged.items():
        for project_id in projects:
            src = find_excel_for_project_set(project_id, set_no)
            if src is None:
                print(f"  skip deploy: project {project_id} set {set_no} (no source xlsx)")
                continue
            dst_dir = Path(f"C:/MOSTest/Excel365/Tab{group}/PracticeVariant{set_no}")
            dst_dir.mkdir(parents=True, exist_ok=True)
            dst = dst_dir / f"project{project_id}.xlsx"
            shutil.copy2(src, dst)
            copied += 1
    print(f"deployed {copied} variant xlsx files to MOSTest")
    return copied


def build_practice_variants_config(merged: dict[int, dict[int, list[dict]]]) -> dict:
    group = "1"
    practice: dict = {group: {}}
    project_ids = set()
    for set_no, projects in merged.items():
        project_ids.update(projects.keys())

    for project_id in sorted(project_ids):
        sets_for_project: dict = {}
        for set_no in range(1, 6):
            if project_id not in merged.get(set_no, {}):
                continue
            excel_path = (
                f"C:\\MOSTest\\Excel365\\Tab{group}\\PracticeVariant{set_no}\\project{project_id}.xlsx"
            )
            if not Path(excel_path).is_file():
                print(f"  skip config: project {project_id} set {set_no} (missing {excel_path})")
                continue
            base_lib = TEXTBOOK_LIBRARY.get(project_id, f"ExcelChecker1_{project_id}")
            sets_for_project[str(set_no)] = {
                "excelFile": excel_path,
                "library": f"{base_lib}_PV{set_no}",
                "tasksJson": f"MOS演習問題文一覧_PracticeVariant{set_no}.json",
            }
        if sets_for_project:
            practice[group][str(project_id)] = sets_for_project
    return practice


def update_config_json(practice_variants: dict) -> None:
    with CONFIG_PATH.open(encoding="utf-8") as f:
        config = json.load(f)
    config["practiceVariants"] = practice_variants
    with CONFIG_PATH.open("w", encoding="utf-8") as f:
        json.dump(config, f, ensure_ascii=False, indent=2)
        f.write("\n")
    print(f"updated {CONFIG_PATH}")


def main() -> int:
    qa_files = sorted(RUIDAI_DIR.glob("完成 Project*_問題解答.xlsx"))
    if not qa_files:
        print(f"no QA workbooks in {RUIDAI_DIR}", file=sys.stderr)
        return 1

    print(f"reading {len(qa_files)} QA workbooks from {RUIDAI_DIR}")
    all_data = []
    for path in qa_files:
        print(f"  {path.name}")
        all_data.append(read_qa_workbook(path))

    merged = merge_variant_data(all_data)
    write_variant_json_files(merged)
    write_answer_csv(merged, qa_files)
    sync_excel_files_to_mostest(merged)
    practice = build_practice_variants_config(merged)
    update_config_json(practice)

    configured_projects = sorted(practice.get("1", {}).keys(), key=int)
    print(f"practiceVariants projects: {', '.join(configured_projects)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
